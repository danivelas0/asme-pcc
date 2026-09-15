# test_motor_declarado.py
import pytest
import motor_declarado as M


def _motor_minimo():
    """Dos filas en una seccion: una entrada y un calculo que la usa."""
    return M.Motor(
        articulo="999",
        hoja="Prueba_PCC2_Art999",
        titulo="MOTOR DE PRUEBA",
        fuente="ASME PCC/pcc_2/p2_welded_repairs/art_999/art_999.json",
        secciones=(
            M.Seccion("1. DATOS DE ENTRADA", filas=(
                M.Fila("P", "Presion de diseno", magnitud="pres", tipo=M.ENTRADA,
                       ejemplo=20,
                       cita=M.Cita("art_999.json", bloque=12, clausula="999-3.2")),
                M.Fila("dos_P", "El doble", magnitud="pres", tipo=M.FORMULA,
                       formula="={P}*2",
                       cita=M.Cita("art_999.json", bloque=12, clausula="999-3.2")),
            )),
        ),
    )


def test_cada_fila_recibe_una_direccion_en_orden():
    dirs = M.resolver_direcciones(_motor_minimo())
    assert dirs["P"] == "$D$6"
    assert dirs["dos_P"] == "$D$7"


def test_la_formula_se_escribe_por_nombre_y_sale_por_direccion():
    dirs = M.resolver_direcciones(_motor_minimo())
    assert M.sustituir_nombres("={P}*2", dirs) == "=$D$6*2"


def test_un_nombre_que_no_existe_aborta():
    dirs = M.resolver_direcciones(_motor_minimo())
    with pytest.raises(SystemExit, match="no existe"):
        M.sustituir_nombres("={NO_EXISTE}+1", dirs)


def test_una_fila_de_calculo_sin_cita_aborta(tmp_path):
    motor = _motor_minimo()._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("x", "Sin procedencia", tipo=M.FORMULA, formula="=1+1"),
        )),))
    with pytest.raises(SystemExit, match="sin procedencia"):
        M.comprobar_procedencia(motor, tmp_path)


def _motor_sin_formulas():
    """Solo una entrada, sin ninguna Fila de tipo FORMULA -y sin pasos-, para
    las pruebas de comprobar_procedencia() que no tienen que abortar: sin una
    FORMULA que citar, la funcion no toca ningun JSON de resources/ y por eso
    estas pruebas no necesitan crear uno en tmp_path."""
    return M.Motor(
        articulo="999",
        hoja="Prueba_PCC2_Art999",
        titulo="MOTOR DE PRUEBA",
        fuente="ASME PCC/pcc_2/p2_welded_repairs/art_999/art_999.json",
        secciones=(
            M.Seccion("1. DATOS DE ENTRADA", filas=(
                M.Fila("P", "Presion de diseno", tipo=M.ENTRADA, ejemplo=20),
            )),
        ),
    )


def test_una_fila_de_calculo_sin_cita_en_un_paso_aborta(tmp_path):
    """Simetrica de test_una_fila_de_calculo_sin_cita_aborta, pero con el
    calculo sin procedencia dentro de un Paso en vez de una Seccion.

    Ronda de arreglo 1 (Hallazgo 1) de la Tarea 6: comprobar_procedencia() se
    amplio para recorrer tambien motor.pasos -via _bloques()-, pero esa
    ampliacion no tenia ninguna prueba que la ejerciera. Sin esta prueba, una
    refactorizacion futura podia dejar de mirar motor.pasos aqui y nada lo
    notaria: es justo el guardia que sostiene la Regla n.1 en el anexo del
    flujo, donde un paso puede calcular algo tan normativo como cualquier
    fila de seccion."""
    motor = _motor_sin_formulas()._replace(pasos=(
        M.Paso(1, "PASO SIN PROCEDENCIA", "999-1", filas=(
            M.Fila("x", "Sin procedencia", tipo=M.FORMULA, formula="=1+1"),
        )),
    ))
    with pytest.raises(SystemExit, match="sin procedencia"):
        M.comprobar_procedencia(motor, tmp_path)


def test_un_motor_sin_pasos_no_aborta_por_procedencia(tmp_path):
    """La otra direccion del Hallazgo 1: un motor sin anexo de pasos
    (motor.pasos == (), el valor por defecto) no le pide procedencia a nada
    que no exista -_bloques() no aporta ningun bloque extra- y
    comprobar_procedencia() no aborta."""
    assert M.comprobar_procedencia(_motor_sin_formulas(), tmp_path) == []


def test_un_paso_con_solo_entradas_no_aborta_por_procedencia(tmp_path):
    """Y la otra mitad de esa misma direccion: un Paso cuyas filas son todas
    ENTRADA -se teclean, no se calculan- tampoco le exige cita a nada. El
    guardia de comprobar_procedencia() solo mira las filas de tipo FORMULA,
    dentro de una Seccion o de un Paso por igual."""
    motor = _motor_sin_formulas()._replace(pasos=(
        M.Paso(1, "PASO SIN CALCULO", "999-1", filas=(
            M.Fila("q", "Se teclea, no se calcula", tipo=M.ENTRADA, ejemplo=5),
        )),
    ))
    assert M.comprobar_procedencia(motor, tmp_path) == []


def test_una_cita_a_un_bloque_que_no_existe_aborta(tmp_path):
    import json
    d = tmp_path / "ASME PCC" / "pcc_2" / "p2_welded_repairs" / "art_999"
    d.mkdir(parents=True)
    (d / "art_999.json").write_text(
        json.dumps({"blocks": [{"type": "paragraph", "text": "uno"}]}),
        encoding="utf-8")
    motor = _motor_minimo()   # su cita apunta al bloque 12, que no existe
    with pytest.raises(SystemExit, match="bloque 12"):
        M.comprobar_procedencia(motor, tmp_path)


def _helpers_de_prueba():
    from openpyxl.styles import Protection

    def banda(ws, fila, texto):
        ws.cell(fila, 1, texto)

    def rotulo(ws, fila, texto, simbolo, unidad, referencia="", comentario=""):
        ws.cell(fila, 1, texto)
        if referencia:
            ws.cell(fila, 7, referencia)

    def entrada(ws, celda, valor):
        ws[celda] = valor
        ws[celda].protection = Protection(locked=False)

    def calculo(ws, celda, formula):
        ws[celda] = formula

    return M.Helpers(banda=banda, rotulo=rotulo, entrada=entrada, calculo=calculo)


def test_la_hoja_sale_con_las_bandas_y_las_filas_en_orden():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    res = M.emitir_tabla(ws, _motor_minimo(), _helpers_de_prueba())
    assert ws["A5"].value == "1. DATOS DE ENTRADA"
    assert ws["A6"].value == "Presion de diseno"
    assert ws["D6"].value == 20                 # el ejemplo del caso precargado
    assert ws["A7"].value == "El doble"
    assert ws["D7"].value == "=$D$6*2"          # la formula, ya con direcciones
    assert res["ultima_fila"] == 7


def test_la_clausula_llega_a_la_columna_de_referencia():
    """La procedencia no se queda en la tabla de trazabilidad: el ingeniero
    tiene que ver de que clausula sale el numero SIN salir de la fila."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    M.emitir_tabla(ws, _motor_minimo(), _helpers_de_prueba())
    assert ws["G6"].value == "999-3.2"


def test_la_entrada_nace_desbloqueada_y_el_calculo_no():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    M.emitir_tabla(ws, _motor_minimo(), _helpers_de_prueba())
    assert ws["D6"].protection.locked is False   # entrada
    assert ws["D7"].protection.locked is not False


def _motor_dos_secciones():
    """Dos secciones con varias filas cada una -no una sola, como en
    `_motor_minimo`-, para blindar que `resolver_direcciones()` y
    `emitir_tabla()` recorren el MISMO arbol. Hoy coinciden por casualidad:
    ninguna prueba anterior lo comprueba con mas de una seccion, que es
    exactamente el patron de "dos fuentes de verdad" que este proyecto marca
    como peligroso en otras partes (ver la capa de navegacion en el CLAUDE.md
    del repo)."""
    cita = M.Cita("art_999.json", bloque=12, clausula="999-3.2")
    return M.Motor(
        articulo="999",
        hoja="Prueba_PCC2_Art999",
        titulo="MOTOR DE PRUEBA",
        fuente="ASME PCC/pcc_2/p2_welded_repairs/art_999/art_999.json",
        secciones=(
            M.Seccion("1. DATOS DE ENTRADA", filas=(
                M.Fila("P", "Presion", tipo=M.ENTRADA, ejemplo=20),
                M.Fila("D", "Diametro", tipo=M.ENTRADA, ejemplo=300),
                M.Fila("dosP", "El doble de P", tipo=M.FORMULA,
                       formula="={P}*2", cita=cita),
            )),
            M.Seccion("2. RESULTADOS", filas=(
                M.Fila("suma", "P + D", tipo=M.FORMULA,
                       formula="={P}+{D}", cita=cita),
                M.Fila("otra", "suma + P", tipo=M.FORMULA,
                       formula="={suma}+{P}", cita=cita),
            )),
        ),
    )


def test_emitir_tabla_escribe_en_la_direccion_que_resolver_direcciones_devuelve():
    """Con dos secciones y varias filas: la direccion que `emitir_tabla` USA
    para escribir cada fila tiene que ser exactamente la que
    `resolver_direcciones` DEVUELVE para esa clave -incluida una formula de la
    segunda seccion que cita una clave de la primera-. Sin esta prueba, las
    dos funciones podian divergir en cuanto un motor real tuviera mas de una
    seccion y nadie lo notaria hasta ver un #REF! en una celda."""
    import openpyxl
    motor = _motor_dos_secciones()
    dirs = M.resolver_direcciones(motor)                      # fuente nº 1

    wb = openpyxl.Workbook()
    ws = wb.active
    M.emitir_tabla(ws, motor, _helpers_de_prueba())            # fuente nº 2

    # Las direcciones en si: la segunda seccion sigue contando filas donde la
    # primera las dejo, y no reinicia en su propia banda.
    assert dirs == {"P": "$D$6", "D": "$D$7", "dosP": "$D$8",
                    "suma": "$D$10", "otra": "$D$11"}

    esperado = {"P": 20, "D": 300, "dosP": "=$D$6*2",
                "suma": "=$D$6+$D$7", "otra": "=$D$10+$D$6"}
    for clave, valor in esperado.items():
        direccion = dirs[clave]
        assert ws[direccion].value == valor, (
            f"{clave}: resolver_direcciones dice {direccion}, pero ese no es "
            f"el valor que emitir_tabla escribio ahi")


def _motor_con_pasos():
    """Una seccion normal seguida del anexo de pasos: dos Paso, cada uno con
    sus propias filas -incluida una formula que cita una clave de la seccion
    y otra que cita una clave de un paso anterior-, para blindar que
    resolver_direcciones() y emitir_tabla() tambien coinciden cuando el
    arbol incluye motor.pasos (Tarea 6 de la skill motor_pcc2)."""
    cita = M.Cita("art_999.json", bloque=12, clausula="999-3.2")
    return M.Motor(
        articulo="999",
        hoja="Prueba_PCC2_Art999",
        titulo="MOTOR DE PRUEBA",
        fuente="ASME PCC/pcc_2/p2_welded_repairs/art_999/art_999.json",
        secciones=(
            M.Seccion("1. DATOS DE ENTRADA", filas=(
                M.Fila("P", "Presion", tipo=M.ENTRADA, ejemplo=20),
            )),
        ),
        pasos=(
            M.Paso(1, "ELEGIBILIDAD", "999-1", filas=(
                M.Fila("elegible", "Compuerta", tipo=M.FORMULA,
                       formula="={P}*0", cita=cita),
            )),
            M.Paso(2, "CARGAS", "999-2", filas=(
                M.Fila("carga", "Carga total", tipo=M.FORMULA,
                       formula="={P}+{elegible}", cita=cita),
            )),
        ),
    )


def test_emitir_tabla_escribe_en_la_direccion_que_resolver_direcciones_devuelve_con_pasos():
    """Mismo blindaje que la prueba anterior, pero con motor.pasos: el anexo
    de pasos tiene que salir DETRAS de la ultima seccion, con una banda por
    paso, y las direcciones que asigna resolver_direcciones() tienen que ser
    las mismas donde emitir_tabla() escribio de verdad -incluida la formula
    del Paso 2 que cita una clave declarada en el Paso 1."""
    import openpyxl
    motor = _motor_con_pasos()
    dirs = M.resolver_direcciones(motor)                       # fuente n.1

    wb = openpyxl.Workbook()
    ws = wb.active
    M.emitir_tabla(ws, motor, _helpers_de_prueba())             # fuente n.2

    # Fila 5: banda de la seccion. Fila 6: "P". Fila 7: banda del Paso 1.
    # Fila 8: "elegible". Fila 9: banda del Paso 2. Fila 10: "carga".
    assert dirs == {"P": "$D$6", "elegible": "$D$8", "carga": "$D$10"}
    assert ws["A7"].value == "PASO 1 · ELEGIBILIDAD — 999-1"
    assert ws["A9"].value == "PASO 2 · CARGAS — 999-2"

    esperado = {"P": 20, "elegible": "=$D$6*0", "carga": "=$D$6+$D$8"}
    for clave, valor in esperado.items():
        direccion = dirs[clave]
        assert ws[direccion].value == valor, (
            f"{clave}: resolver_direcciones dice {direccion}, pero ese no es "
            f"el valor que emitir_tabla escribio ahi")


# ---------------------------------------------------------------------------
# Ronda de arreglo final de la rama skill_motor_pcc2 (2026-09-13)
# ---------------------------------------------------------------------------
def _json_de_prueba(tmp_path, blocks):
    """Escribe el art_999.json que citan los motores de prueba."""
    import json
    d = tmp_path / "ASME PCC" / "pcc_2" / "p2_welded_repairs" / "art_999"
    d.mkdir(parents=True, exist_ok=True)
    (d / "art_999.json").write_text(json.dumps({"blocks": blocks}),
                                    encoding="utf-8")
    return tmp_path


def test_citar_un_section_header_aborta(tmp_path):
    """Un ROTULO no publica un valor. El caso real que lo destapo:
    `gap_ok` de la plantilla citaba el bloque 62 de art_206.json, que es el
    section_header "206-4.1 Installation"; la luz radial de 2,5 mm la imprime
    el parrafo 63. El bloque EXISTE, asi que el guardia de rango pasaba y la
    cita parecia valida apuntando a donde el dato no esta."""
    _json_de_prueba(tmp_path, [
        {"type": "paragraph", "text": "cero"},
        {"type": "section_header", "text": "999-3.2 Un rotulo"},
    ])
    motor = _motor_minimo()
    motor = motor._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("x", "Cita a un rotulo", tipo=M.FORMULA, formula="=1+1",
                   cita=M.Cita("art_999.json", bloque=1, clausula="999-3.2")),
        )),))
    with pytest.raises(SystemExit, match="section_header"):
        M.comprobar_procedencia(motor, tmp_path)


def test_citar_el_parrafo_que_si_publica_el_dato_no_aborta(tmp_path):
    """La otra direccion: el guardia rechaza el rotulo, no la cita correcta."""
    _json_de_prueba(tmp_path, [
        {"type": "section_header", "text": "999-3.2 Un rotulo"},
        {"type": "paragraph", "text": "El valor es 2.5 mm"},
    ])
    motor = _motor_minimo()._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("x", "Cita al parrafo", tipo=M.FORMULA, formula="=1+1",
                   cita=M.Cita("art_999.json", bloque=1, clausula="999-3.2")),
        )),))
    assert M.comprobar_procedencia(motor, tmp_path) == [
        ("x", "999-3.2", "art_999.json", 1)]


def test_verificar_py_pasa_a_comprobar_procedencia_un_tipo_que_acepta(tmp_path):
    """La §12 de verificar.py llamaba `MD.comprobar_procedencia(motor, RES)`,
    y `RES` es un `db_lib.Resources`, que NO es os.PathLike: `Path(resources)`
    reventaba con TypeError dentro del try/finally que cierra Excel, asi que
    `auditar()` moria sin escribir el reporte. Con el registro de motores
    vacio el bucle no itera nunca, asi que ninguna corrida lo destapaba.

    Esta prueba no reimplementa la llamada: LEE la expresion real del segundo
    argumento en el fuente de verificar.py (via ast), la evalua contra un
    `Resources` de verdad y con el resultado ejerce `comprobar_procedencia`.
    Si alguien vuelve a pasar `RES` -o cualquier objeto que la funcion no
    acepte-, esta prueba falla sin necesidad de Excel ni de un motor real."""
    import ast
    import os
    from pathlib import Path

    import db_lib

    fuente = (Path(__file__).resolve().parent / "verificar.py").read_text(
        encoding="utf-8")
    llamadas = [
        n for n in ast.walk(ast.parse(fuente))
        if isinstance(n, ast.Call)
        and isinstance(n.func, ast.Attribute)
        and n.func.attr == "comprobar_procedencia"]
    assert len(llamadas) == 1, "se esperaba UNA llamada en verificar.py"
    expr = ast.unparse(llamadas[0].args[1])

    res = db_lib.Resources(str(tmp_path))
    valor = eval(expr, {"RES": res})              # noqa: S307 - expresion del repo
    assert isinstance(valor, (str, os.PathLike)), (
        f"verificar.py pasa {expr!r} = {type(valor).__name__}, que "
        f"comprobar_procedencia no puede convertir en Path")
    # Y se ejerce de verdad, con ese mismo valor.
    assert M.comprobar_procedencia(_motor_sin_formulas(), valor) == []


def _motor_con_dictamen():
    """Motor con las cuatro claves reservadas y una fila de dictamen."""
    cita = M.Cita("art_999.json", bloque=0, clausula="999-3.2")
    return M.Motor(
        articulo="999", hoja="Prueba_PCC2_Art999", titulo="MOTOR DE PRUEBA",
        fuente="ASME PCC/pcc_2/p2_welded_repairs/art_999/art_999.json",
        secciones=(
            M.Seccion("1. APLICACION", filas=(
                M.Fila("unidad", "Sistema de unidades", tipo=M.LISTA,
                       ejemplo="SI", lista=M.Lista(opciones=("SI", "US"))),
                M.Fila("modo", "Codigo de aplicacion", tipo=M.LISTA,
                       ejemplo="PCC2-999",
                       lista=M.Lista(opciones=("PCC2-999",))),
                M.Fila("temperatura", "Temperatura", magnitud="temp",
                       tipo=M.ENTRADA, ejemplo=20),
            )),
            M.Seccion("2. RESULTADO", filas=(
                M.Fila("chk", "Verificacion", tipo=M.FORMULA,
                       formula='=IF({temperatura}>0,"CUMPLE","NO CUMPLE")',
                       cita=cita),
                # Sin formula: la compone el chasis desde Motor.dictamen.
                M.Fila("dictamen", "Dictamen global", tipo=M.FORMULA,
                       cita=cita),
            )),
        ),
        verificaciones=(M.Verificacion(
            clave="chk", rotulo="Verificacion", requerido="{temperatura}",
            adoptado="{temperatura}", criterio="T > 0"),),
        dictamen=M.Dictamen(verificaciones=("chk",)),
    )


def test_un_nombre_inexistente_en_una_verificacion_aborta():
    motor = _motor_con_dictamen()
    motor = motor._replace(verificaciones=(
        motor.verificaciones[0]._replace(requerido="{no_existe}"),))
    with pytest.raises(SystemExit, match="no_existe"):
        M.comprobar_nombres_declarativos(motor)


def test_los_nombres_de_una_verificacion_valida_no_abortan():
    assert M.comprobar_nombres_declarativos(_motor_con_dictamen()) is None


def test_el_texto_de_una_especificacion_sustituye_sus_nombres():
    """`Especificacion.texto` se documenta como "formula con {nombres}" y se
    copiaba CRUDO a la pestana: un `{t_req}` salia impreso con las llaves. La
    direccion va CALIFICADA CON LA HOJA del motor -la pestana es otra hoja y
    un `$D$9` a secas apuntaria dentro de ella-."""
    motor = _motor_con_dictamen()
    e = M.Especificacion(concepto="Dictamen", texto="={dictamen}",
                         clausula="999-6")
    assert M.texto_de_especificacion(motor, e) == \
        "='Prueba_PCC2_Art999'!$D$11"


def test_un_nombre_inexistente_en_una_especificacion_aborta():
    e = M.Especificacion(concepto="X", texto="={no_existe}", clausula="999-6")
    with pytest.raises(SystemExit, match="no existe"):
        M.texto_de_especificacion(_motor_con_dictamen(), e)


def test_un_texto_fijo_de_especificacion_pasa_intacto():
    e = M.Especificacion(concepto="X", texto="100 mm (4 in.)", clausula="999-6")
    assert M.texto_de_especificacion(_motor_con_dictamen(), e) == \
        "100 mm (4 in.)"


def test_bloques_es_el_alias_publico_del_recorrido():
    """Tres consumidores del builder recorrian `motor.secciones` por su cuenta
    e ignoraban `motor.pasos`. El alias existe para que no haya excusa para
    escribir un cuarto recorrido paralelo."""
    assert M.bloques is M._bloques
    titulos = [t for t, _f in M.bloques(_motor_con_pasos())]
    assert titulos == ["1. DATOS DE ENTRADA",
                       "PASO 1 · ELEGIBILIDAD — 999-1",
                       "PASO 2 · CARGAS — 999-2"]


# ---------------------------------------------------------------------------
# tipo=LISTA atada a una base de datos (instruccion del ingeniero 2026-09-13:
# "todo material debe ser extraido de las bases de datos existentes, siempre")
# ---------------------------------------------------------------------------
def _motor_con_lista(lista):
    """Un motor de una sola fila LISTA, para ejercer los guardias."""
    return M.Motor(
        articulo="999", hoja="Prueba_PCC2_Art999", titulo="MOTOR DE PRUEBA",
        fuente="ASME PCC/pcc_2/p2_welded_repairs/art_999/art_999.json",
        secciones=(M.Seccion("1. DATOS", filas=(
            M.Fila("NPS", "Diametro nominal", tipo=M.LISTA, lista=lista),
        )),),
    )


def test_una_fila_lista_sin_base_aborta():
    """Es el guardia central de la instruccion: una lista que no dice de que
    base sale es una lista tecleada a mano, y esas no se auditan."""
    with pytest.raises(SystemExit, match="de que base sale"):
        M.comprobar_listas(_motor_con_lista(None))


def test_una_fila_lista_atada_a_una_base_pasa():
    assert M.comprobar_listas(_motor_con_lista(
        M.Lista(hoja="DB_B36_10", columna="NPS impreso"))) is None


def test_una_enumeracion_tecleada_fuera_de_las_claves_reservadas_aborta():
    """La enumeracion solo se admite en las claves del propio marco. Fuera de
    ellas es exactamente la lista fija que la regla 12 prohibe."""
    with pytest.raises(SystemExit, match="enumeracion tecleada"):
        M.comprobar_listas(_motor_con_lista(M.Lista(opciones=("2", "3", "4"))))


def test_una_enumeracion_en_una_clave_reservada_pasa():
    motor = _motor_con_lista(None)._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("unidad", "Sistema de unidades", tipo=M.LISTA,
                   lista=M.Lista(opciones=("SI", "US"))),)),))
    assert M.comprobar_listas(motor) is None


def test_una_lista_con_base_y_enumeracion_a_la_vez_aborta():
    with pytest.raises(SystemExit, match="base y una enumeracion"):
        M.comprobar_listas(_motor_con_lista(
            M.Lista(hoja="DB_B36_10", columna="NPS impreso",
                    opciones=("1", "2"))))


def test_una_lista_declarada_en_una_fila_que_no_es_lista_aborta():
    motor = _motor_con_lista(None)._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("x", "Entrada", tipo=M.ENTRADA,
                   lista=M.Lista(hoja="DB_B36_10", columna="NPS impreso")),)),))
    with pytest.raises(SystemExit, match="su tipo es"):
        M.comprobar_listas(motor)


def test_emitir_tabla_sin_helper_de_lista_aborta():
    """Sin el helper, la celda quedaria editable a mano y sin desplegable: la
    regla 12 incumplida en silencio, que es el estado anterior a este cambio."""
    import openpyxl
    ws = openpyxl.Workbook().active
    motor = _motor_con_lista(M.Lista(hoja="DB_B36_10", columna="NPS impreso"))
    with pytest.raises(SystemExit, match="no traen"):
        M.emitir_tabla(ws, motor, _helpers_de_prueba())


def test_emitir_tabla_ejerce_los_guardias_del_marco():
    """Los tres guardias corren dentro de `emitir_tabla`, que es la unica
    puerta por la que una declaracion se convierte en hoja. Probarlos solo
    llamandolos a mano no demuestra que esten ENCHUFADOS: un guardia que se
    pueda esquivar llamando a otra funcion no es un guardia."""
    import openpyxl
    helpers = _helpers_de_prueba()._replace(lista=lambda w, c, d: None)
    # (a) lista sin base
    with pytest.raises(SystemExit, match="de que base sale"):
        M.emitir_tabla(openpyxl.Workbook().active,
                       _motor_con_lista(None), helpers)
    # (b) dos casos
    with pytest.raises(SystemExit, match="UNA columna de valor"):
        M.emitir_tabla(openpyxl.Workbook().active,
                       _motor_minimo()._replace(casos=("A", "B")), helpers)
    # (c) fila de dictamen con formula tecleada
    motor = _motor_con_dictamen()
    motor = motor._replace(secciones=tuple(
        s._replace(filas=tuple(
            f._replace(formula='=IF(1=1,"APTO","REVISAR")')
            if f.clave == "dictamen" else f for f in s.filas))
        for s in motor.secciones))
    with pytest.raises(SystemExit, match="formula tecleada"):
        M.emitir_tabla(openpyxl.Workbook().active, motor, helpers)


def test_emitir_tabla_llama_al_helper_de_lista_con_su_declaracion():
    import openpyxl
    ws = openpyxl.Workbook().active
    decl = M.Lista(hoja="DB_B36_10", columna="NPS impreso")
    vistas = []
    helpers = _helpers_de_prueba()._replace(
        lista=lambda w, celda, d: vistas.append((celda, d)))
    M.emitir_tabla(ws, _motor_con_lista(decl), helpers)
    assert vistas == [("D6", decl)]


# ---------------------------------------------------------------------------
# casos: UNA columna de valor, y el chasis lo dice en vez de fingirlo
# ---------------------------------------------------------------------------
def test_declarar_dos_casos_aborta():
    motor = _motor_minimo()._replace(casos=("Operacion", "Diseno"))
    with pytest.raises(SystemExit, match="UNA columna de valor"):
        M.comprobar_casos(motor)


def test_un_solo_caso_pasa():
    assert M.comprobar_casos(_motor_minimo()) is None


def test_el_defecto_de_casos_es_uno_solo():
    """El defecto prometia dos columnas y el chasis emitia una: la segunda
    salia vacia y sin aviso."""
    assert len(_motor_minimo().casos) == 1


# ---------------------------------------------------------------------------
# El dictamen lo COMPONE el chasis
# ---------------------------------------------------------------------------
def test_el_chasis_compone_el_dictamen_desde_lo_declarado():
    motor = _motor_con_dictamen()
    dirs = M.resolver_direcciones(motor)
    assert M.formula_dictamen(motor) == (
        '=IF(AND(' + dirs["chk"] + '="CUMPLE"),"APTO","REVISAR")')


def test_el_dictamen_compuesto_llega_a_la_hoja():
    import openpyxl
    ws = openpyxl.Workbook().active
    motor = _motor_con_dictamen()
    helpers = _helpers_de_prueba()._replace(lista=lambda w, c, d: None)
    M.emitir_tabla(ws, motor, helpers)
    celda = M.resolver_direcciones(motor)["dictamen"].replace("$", "")
    assert ws[celda].value == M.formula_dictamen(motor)


def test_las_compuertas_envuelven_al_and_y_publican_su_propio_texto():
    """Una compuerta que bloquea pasa a SER el dictamen: un motivo dice mas
    que un REVISAR pelado."""
    motor = _motor_con_dictamen()
    motor = motor._replace(
        verificaciones=motor.verificaciones + (M.Verificacion(
            clave="temperatura", rotulo="Compuerta", requerido="", adoptado="",
            criterio="", favorables=("ELEGIBLE",), avisos=("REVISAR",)),),
        dictamen=M.Dictamen(compuertas=("temperatura",),
                            verificaciones=("chk",)))
    g = M.resolver_direcciones(motor)["temperatura"]
    f = M.formula_dictamen(motor)
    assert f.startswith(
        '=IF(NOT(OR(' + g + '="ELEGIBLE",LEFT(' + g + ',7)="REVISAR")),'
        + g + ',')
    assert '"APTO","REVISAR")' in f


def test_el_material_se_antepone_al_and_de_verificaciones():
    motor = _motor_con_dictamen()
    f = M.formula_dictamen(motor, celdas_material=("$D$40",))
    assert '$D$40="' + M.TXT_SIN_MATERIAL + '"' in f
    assert '$D$40<>"' + M.TXT_MATERIAL_OK + '"' in f
    # Y el AND sigue dentro, no sustituido.
    assert '"APTO","REVISAR")' in f


def test_una_formula_de_dictamen_tecleada_aborta():
    """Teclearla al lado de la declaracion es la forma segura de que las dos
    digan cosas distintas."""
    motor = _motor_con_dictamen()
    secciones = tuple(
        s._replace(filas=tuple(
            f._replace(formula='=IF(1=1,"APTO","REVISAR")')
            if f.clave == "dictamen" else f for f in s.filas))
        for s in motor.secciones)
    with pytest.raises(SystemExit, match="formula tecleada"):
        M.comprobar_dictamen(motor._replace(secciones=secciones))


def test_una_fila_dictamen_sin_declaracion_aborta():
    with pytest.raises(SystemExit, match="van juntos"):
        M.comprobar_dictamen(_motor_con_dictamen()._replace(dictamen=None))


def test_una_verificacion_del_and_con_otro_favorable_aborta():
    """verificar.py s.12 juzga el dictamen contra el literal CUMPLE: con otro
    favorable, el guardia mediria una cosa distinta de la que el motor
    calcula."""
    motor = _motor_con_dictamen()
    motor = motor._replace(verificaciones=(
        motor.verificaciones[0]._replace(favorables=("APTO",)),))
    with pytest.raises(SystemExit, match="entra al AND"):
        M.comprobar_dictamen(motor)


def test_una_compuerta_sin_verificacion_declarada_aborta():
    motor = _motor_con_dictamen()._replace(
        dictamen=M.Dictamen(compuertas=("temperatura",),
                            verificaciones=("chk",)))
    with pytest.raises(SystemExit, match="no esta declarada como"):
        M.comprobar_dictamen(motor)


# ---------------------------------------------------------------------------
# El dictamen no se elige a mano (revision del bloque plantilla, 2026-09-13)
# ---------------------------------------------------------------------------
def _con_dictamen_de_lista():
    """El mismo motor, con la fila `dictamen` convertida en desplegable.

    Es la declaracion que el revisor ejecuto y que los TRES guardias dejaban
    pasar: `comprobar_dictamen` solo miraba `fila.formula` -vacia en una fila
    LISTA- y `emitir_tabla` solo componia la formula en la rama FORMULA. La
    celda del veredicto global quedaba editable, amarilla, con desplegable y
    dentro del manifiesto de reinicio.
    """
    motor = _motor_con_dictamen()
    secciones = tuple(
        s._replace(filas=tuple(
            f._replace(tipo=M.LISTA, ejemplo="APTO",
                       lista=M.Lista(opciones=("APTO", "REVISAR")))
            if f.clave == "dictamen" else f for f in s.filas))
        for s in motor.secciones)
    return motor._replace(secciones=secciones)


def test_la_fila_dictamen_declarada_como_desplegable_aborta():
    """Primera via: el tipo. El veredicto global lo CALCULA el libro."""
    with pytest.raises(SystemExit, match="tipo='lista'"):
        M.comprobar_dictamen(_con_dictamen_de_lista())


def test_la_fila_dictamen_declarada_como_entrada_aborta():
    """Sin `Lista` de por medio -tipo=ENTRADA- la fila tampoco trae formula,
    asi que el guardia viejo tampoco la veia."""
    motor = _motor_con_dictamen()
    secciones = tuple(
        s._replace(filas=tuple(
            f._replace(tipo=M.ENTRADA, ejemplo="APTO")
            if f.clave == "dictamen" else f for f in s.filas))
        for s in motor.secciones)
    with pytest.raises(SystemExit, match="tipo='entrada'"):
        M.comprobar_dictamen(motor._replace(secciones=secciones))


def test_emitir_tabla_rechaza_el_dictamen_de_desplegable():
    """Segunda via, y por la puerta real: `emitir_tabla` es lo unico que
    convierte una declaracion en hoja.

    Los helpers llevan `lista` a proposito. Sin el, esta prueba pasaba
    revirtiendo LOS DOS guardias -abortaba con "no traen `lista`", que es otra
    cosa-: una prueba que se pone verde por el motivo equivocado es justo lo
    que esta rama ya ha pagado cuatro veces. Con el helper puesto, el unico
    motivo posible de aborto es el guardia del dictamen.
    """
    import openpyxl
    h = _helpers_de_prueba()._replace(lista=lambda ws, celda, decl: None)
    ws = openpyxl.Workbook().active
    with pytest.raises(SystemExit, match="dictamen"):
        M.emitir_tabla(ws, _con_dictamen_de_lista(), h)
    # Y la declaracion legitima -misma fila, tipo FORMULA- SI construye: el
    # guardia rechaza el desplegable, no el motor.
    ws2 = openpyxl.Workbook().active
    M.emitir_tabla(ws2, _motor_con_dictamen(), h)


def test_la_enumeracion_tecleada_solo_vale_en_unidad_y_modo():
    """La puerta de `Lista(opciones=...)` NO son las cuatro claves
    reservadas: `temperatura` es un dato de servicio que se teclea y
    `dictamen` es el veredicto que compone el chasis. Solo el conmutador
    SI/US y el selector de codigo son enumeraciones del marco."""
    assert M.CLAVES_CON_ENUMERACION == ("unidad", "modo")
    for clave in ("temperatura", "dictamen"):
        motor = _motor_con_lista(None)._replace(secciones=(
            M.Seccion("1. DATOS", filas=(
                M.Fila(clave, "Reservada", tipo=M.LISTA,
                       lista=M.Lista(opciones=("A", "B"))),)),))
        with pytest.raises(SystemExit, match="enumeracion tecleada"):
            M.comprobar_listas(motor)


def test_el_ejemplo_de_una_enumeracion_tiene_que_estar_entre_sus_items():
    """Una validacion 'detener' no rechaza un valor que ya estaba puesto: un
    ejemplo fuera de lista se queda con aspecto de dato valido."""
    motor = _motor_con_lista(None)._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("unidad", "Sistema de unidades", tipo=M.LISTA,
                   ejemplo="Metrico", lista=M.Lista(opciones=("SI", "US"))),)),))
    with pytest.raises(SystemExit, match="no esta entre los items"):
        M.comprobar_listas(motor)


def test_el_ejemplo_que_si_esta_entre_los_items_pasa():
    motor = _motor_con_lista(None)._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("unidad", "Sistema de unidades", tipo=M.LISTA,
                   ejemplo="SI", lista=M.Lista(opciones=("SI", "US"))),)),))
    assert M.comprobar_listas(motor) is None


# ---------------------------------------------------------------------------
# Multifuente (2026-09-15). Un articulo que DELEGA el calculo -el Art. 204 lo
# hace ocho veces- cita el codigo de construccion, no solo su propio JSON.
# ---------------------------------------------------------------------------
from pathlib import Path as _Path

_RESOURCES = _Path(__file__).resolve().parents[3] / "resources"
_B313_CAP2 = "ASME B31/ASME B31.3/CHAPTERS/chapter_02.json"


def _segundo_json(tmp_path, blocks, rel="otro/art_888.json"):
    import json
    ruta = tmp_path / rel
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(json.dumps({"blocks": blocks}), encoding="utf-8")
    return rel


def test_bloques_citables_deja_intacta_la_forma_de_marker():
    """Un JSON de PCC-2 ya viene plano: no se toca."""
    doc = {"blocks": [{"type": "paragraph", "text": "a"}]}
    assert M.bloques_citables(doc) is doc["blocks"]


def test_bloques_citables_aplana_un_capitulo_del_b31_3():
    """Encabezado, luego sus parrafos, luego sus subsecciones: preorden."""
    doc = {"preamble": [{"kind": "paragraph", "text": "pre"}],
           "sections": [{"id": "1", "heading": "Uno",
                         "paragraphs": [{"kind": "paragraph", "text": "p1"},
                                        {"kind": "equation", "text": "e1"}],
                         "sections": [{"id": "1.1", "heading": "Uno uno",
                                       "paragraphs": [{"kind": "paragraph",
                                                       "text": "p11"}]}]}]}
    assert [(b["type"], b["text"]) for b in M.bloques_citables(doc)] == [
        ("paragraph", "pre"),
        ("section_header", "1 Uno"),
        ("paragraph", "p1"),
        ("equation", "e1"),
        ("section_header", "1.1 Uno uno"),
        ("paragraph", "p11"),
    ]


def test_bloques_citables_aplana_una_tabla_suelta():
    """El rotulo primero -y es section_header, asi que no se puede citar- y
    despues una entrada por fila impresa."""
    doc = {"table_id": "Table 304.4.1-1", "title": "Closures",
           "columns": [{"key": "a"}], "rows": [{"a": 1}, {"a": 2}]}
    bs = M.bloques_citables(doc)
    assert bs[0]["type"] == "section_header"
    assert [b["type"] for b in bs[1:]] == ["table_row", "table_row"]
    assert len(bs) == 3


def test_el_aplanado_del_b31_3_es_estable():
    """EL guardia del aplanado, y el motivo de que exista.

    `Cita.bloque` es un INDICE. Si el aplanado cambiara de orden o de criterio,
    TODAS las citas al B31.3 se re-apuntarian a otro parrafo **en silencio** y
    ningun otro guardia lo veria: el bloque seguiria existiendo y seguiria sin
    ser un rotulo. Se fija aqui contra el archivo real: la eq. (3a) de
    `para. 304.1.2` tiene que seguir siendo el bloque 297.
    """
    import json
    doc = json.loads((_RESOURCES / _B313_CAP2).read_text(encoding="utf-8"))
    bs = M.bloques_citables(doc)
    assert len(bs) == 1106
    assert bs[295]["type"] == "section_header"
    assert bs[295]["text"] == "304.1.2 Straight Pipe Under Internal Pressure"
    assert bs[297]["type"] == "equation"
    assert "(3a)" in bs[297]["text"] and "SEW" in bs[297]["text"]
    # Las tres del refuerzo por area, que son las que calculan la abertura
    # mecanizada en el cap de la caja (304.4.2(e) remite a 304.3.3).
    assert "(6)" in bs[401]["text"]
    assert "(7)" in bs[407]["text"]
    assert "(8)" in bs[409]["text"]


def test_una_cita_con_fuente_propia_se_valida_contra_ESE_archivo(tmp_path):
    _json_de_prueba(tmp_path, [{"type": "paragraph", "text": "del 999"}])
    rel = _segundo_json(tmp_path, [{"type": "section_header", "text": "rotulo"},
                                   {"type": "equation", "text": "t = PD/2SE"}])
    motor = _motor_minimo()._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("t", "Espesor delegado", tipo=M.FORMULA, formula="=1",
                   cita=M.Cita("B31.3 cap. II", bloque=1, clausula="304.1.2",
                               fuente=rel)),
        )),))
    assert M.comprobar_procedencia(motor, tmp_path) == [
        ("t", "304.1.2", "B31.3 cap. II", 1)]


def test_una_cita_a_un_segundo_archivo_que_no_existe_aborta(tmp_path):
    _json_de_prueba(tmp_path, [{"type": "paragraph", "text": "x"}])
    motor = _motor_minimo()._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("t", "Espesor", tipo=M.FORMULA, formula="=1",
                   cita=M.Cita("B31.3", bloque=0, clausula="304.1.2",
                               fuente="no/existe.json")),
        )),))
    with pytest.raises(SystemExit, match="no existe"):
        M.comprobar_procedencia(motor, tmp_path)


def test_el_guardia_del_rotulo_tambien_rige_en_el_segundo_archivo(tmp_path):
    _json_de_prueba(tmp_path, [{"type": "paragraph", "text": "x"}])
    rel = _segundo_json(tmp_path, [{"type": "section_header", "text": "304.4.1"}])
    motor = _motor_minimo()._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("t", "Espesor", tipo=M.FORMULA, formula="=1",
                   cita=M.Cita("B31.3", bloque=0, clausula="304.4.1",
                               fuente=rel)),
        )),))
    with pytest.raises(SystemExit, match="section_header"):
        M.comprobar_procedencia(motor, tmp_path)


def test_la_cita_de_una_verificacion_tambien_se_comprueba(tmp_path):
    """Hasta 2026-09-15 nadie las miraba: una cita rota pasaba en verde."""
    _json_de_prueba(tmp_path, [{"type": "paragraph", "text": "x"}])
    motor = _motor_minimo()._replace(
        secciones=(),
        verificaciones=(M.Verificacion(
            "v", "Una verificacion", requerido="", adoptado="", criterio="a>=b",
            cita=M.Cita("art_999.json", bloque=99, clausula="999-5")),))
    with pytest.raises(SystemExit, match="verificacion"):
        M.comprobar_procedencia(motor, tmp_path)


def test_la_cita_de_una_especificacion_tambien_se_comprueba(tmp_path):
    _json_de_prueba(tmp_path, [{"type": "paragraph", "text": "x"}])
    motor = _motor_minimo()._replace(
        secciones=(),
        especificaciones=(("Grupo", (M.Especificacion(
            "Un concepto", "Un texto", "204-4.1",
            cita=M.Cita("art_999.json", bloque=99, clausula="204-4.1")),)),))
    with pytest.raises(SystemExit, match="especificacion"):
        M.comprobar_procedencia(motor, tmp_path)


def test_una_enumeracion_del_codigo_sin_cita_aborta():
    """Regla 12: una lista tecleada que nadie puede auditar no entra."""
    motor = _motor_minimo()._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("metodo_prueba", "Metodo de prueba", tipo=M.LISTA,
                   ejemplo="Hidrostatica",
                   lista=M.Lista(opciones=("Hidrostatica", "Neumatica"))),
        )),))
    with pytest.raises(SystemExit, match="cita"):
        M.comprobar_listas(motor)


def test_una_enumeracion_del_codigo_con_cita_se_admite():
    motor = _motor_minimo()._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("metodo_prueba", "Metodo de prueba", tipo=M.LISTA,
                   ejemplo="Hidrostatica",
                   lista=M.Lista(opciones=("Hidrostatica", "Neumatica"),
                                 cita=M.Cita("art_204.json", bloque=101,
                                             clausula="204-6.2"))),
        )),))
    assert M.comprobar_listas(motor) is None          # no aborta


def test_una_enumeracion_del_MARCO_con_cita_aborta():
    """El conmutador SI/US no lo publica ningun articulo: citarlo seria
    afirmar una procedencia que no existe."""
    motor = _motor_minimo()._replace(secciones=(
        M.Seccion("1. DATOS", filas=(
            M.Fila("unidad", "Sistema de unidades", tipo=M.LISTA, ejemplo="SI",
                   lista=M.Lista(opciones=("SI", "US"),
                                 cita=M.Cita("art_204.json", bloque=2,
                                             clausula="204-1"))),
        )),))
    with pytest.raises(SystemExit, match="MARCO"):
        M.comprobar_listas(motor)


# ---------------------------------------------------------------------------
# {S} de la cascada dentro de una ecuacion (2026-09-15). Lo pide el Art. 204:
# delega el espesor al B31.3 y la eq. (3a) es t = PD/(2(SEW + PY)).
# ---------------------------------------------------------------------------
def _motor_con_S(formula="={P}*2/(2*({S}*{E}))", columnas=(("D", "Metal base"),)):
    return _motor_minimo()._replace(
        material=M.Material(columnas=columnas),
        secciones=(
            M.Seccion("1. DATOS", filas=(
                M.Fila("P", "Presion", magnitud="pres", tipo=M.ENTRADA,
                       ejemplo=20),
                M.Fila("E", "Eficiencia", tipo=M.ENTRADA, ejemplo=1.0),
                M.Fila("t_req", "Espesor requerido", magnitud="len",
                       tipo=M.FORMULA, formula=formula,
                       cita=M.Cita("B31.3", bloque=1, clausula="304.1.2")),
            )),))


def test_el_nombre_de_material_sale_como_centinela():
    """La direccion del S(T) no existe cuando se emite la formula: el bloque
    de material se construye DESPUES de la tabla."""
    salida = M.sustituir_nombres("={S}*2", {}, "D")
    assert salida == "=__MAT@S@D__*2"
    # El centinela lleva '@', que Excel no admite en una referencia: uno que
    # sobreviva a los dos pases revienta en vez de calcular contra otra celda.
    assert "@" in salida


def test_se_puede_pedir_el_material_de_otra_columna():
    assert M.sustituir_nombres("={S@E}", {}, "D") == "=__MAT@S@E__"


def test_nombres_de_material_los_encuentra_en_las_formulas():
    assert M.nombres_de_material(_motor_con_S()) == {("S", "D")}


def test_pedir_S_sin_cascada_de_material_aborta():
    motor = _motor_con_S()._replace(material=None)
    with pytest.raises(SystemExit, match="no declara `material`"):
        M.comprobar_material_citable(motor)


def test_pedir_una_columna_de_material_que_no_existe_aborta():
    motor = _motor_con_S(formula="={P}*{S@E}")
    with pytest.raises(SystemExit, match="columna 'E'"):
        M.comprobar_material_citable(motor)


def test_un_motor_sin_material_en_formulas_no_exige_cascada():
    assert M.comprobar_material_citable(_motor_minimo()) is None


def test_resolver_material_cambia_el_centinela_por_la_direccion():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["D6"] = "=__MAT@S@D__*2"
    ws["D7"] = "=1+1"
    refs = {"por_columna": {"D": {"s_t": "$D$99", "tmax": "$D$97",
                                  "material_id": "$D$88", "dictamen": "$D$98"}}}
    tocadas = M.resolver_material(ws, _motor_con_S(), refs, 10, 7)
    assert tocadas == 1
    assert ws["D6"].value == "=$D$99*2"
    assert ws["D7"].value == "=1+1"
